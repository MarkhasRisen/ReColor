"""
Convert recolor_test_cases_consolidated.md → Word (.docx) with proper
table formatting, section headings, and professional styling.
"""

import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── Paths ──────────────────────────────────────────────────────────────
MD_PATH = r"C:\Users\markr\.gemini\antigravity\brain\09b7a398-5f89-40b9-a86d-6c7983fe12c2\recolor_test_cases_consolidated.md"
DOCX_PATH = r"c:\Users\markr\OneDrive\Desktop\Local Project Filee\ReColor\evaluation\recolor_test_cases_consolidated.docx"

# ── Helpers ────────────────────────────────────────────────────────────

def set_cell_shading(cell, color_hex):
    """Apply background shading to a table cell."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, **kwargs):
    """Set borders on a cell. kwargs: top, bottom, left, right, each a dict with sz, val, color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, attrs in kwargs.items():
        el = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{attrs.get("val","single")}" '
            f'w:sz="{attrs.get("sz","4")}" w:space="0" '
            f'w:color="{attrs.get("color","000000")}"/>'
        )
        tcBorders.append(el)
    tcPr.append(tcBorders)


def style_header_row(row, bg_color="1F3864"):
    """Style a table header row with dark background and white bold text."""
    for cell in row.cells:
        set_cell_shading(cell, bg_color)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(8)
                run.font.name = "Calibri"


def style_data_cell(cell, font_size=Pt(8), align=WD_ALIGN_PARAGRAPH.LEFT):
    """Apply consistent styling to a data cell."""
    for para in cell.paragraphs:
        para.alignment = align
        pf = para.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        for run in para.runs:
            run.font.size = font_size
            run.font.name = "Calibri"


def add_alternating_shading(table):
    """Zebra-stripe data rows."""
    for i, row in enumerate(table.rows):
        if i == 0:
            continue  # header already styled
        if i % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "D6E4F0")


def set_column_widths(table, widths_cm):
    """Set column widths in centimetres."""
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


def set_narrow_margins(doc):
    """Set narrow page margins (1.27 cm / 0.5 in all around) and landscape."""
    for section in doc.sections:
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)
        # Landscape
        new_width = section.page_height
        new_height = section.page_width
        section.page_width = new_width
        section.page_height = new_height


# ── Parse Markdown ─────────────────────────────────────────────────────

def parse_md(path):
    """
    Parse the markdown file into a list of sections. Each section is a dict:
      { "heading": str, "level": int, "blockquote": str|None, "tables": [ [[row], ...] ] }
    """
    with open(path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    sections = []
    current = None

    for line in lines:
        raw = line.rstrip("\n\r")

        # Heading
        m = re.match(r"^(#{1,4})\s+(.+)$", raw)
        if m:
            current = {
                "heading": m.group(2).strip(),
                "level": len(m.group(1)),
                "blockquote": None,
                "tables": [],
                "_in_table": False,
                "_current_table": [],
            }
            sections.append(current)
            continue

        if current is None:
            # blockquote before first heading → attach to a pseudo-section
            if raw.startswith(">"):
                current = {
                    "heading": None,
                    "level": 0,
                    "blockquote": raw.lstrip("> ").strip(),
                    "tables": [],
                    "_in_table": False,
                    "_current_table": [],
                }
                sections.append(current)
            continue

        # Blockquote
        if raw.startswith(">"):
            current["blockquote"] = raw.lstrip("> ").strip()
            continue

        # Horizontal rule
        if raw.strip() == "---":
            continue

        # Table row
        if "|" in raw and raw.strip().startswith("|"):
            cells = [c.strip() for c in raw.strip().strip("|").split("|")]
            # Skip separator rows (---|---|---)
            if all(re.match(r"^[-:]+$", c) for c in cells):
                continue
            if not current["_in_table"]:
                current["_in_table"] = True
                current["_current_table"] = []
            current["_current_table"].append(cells)
        else:
            if current["_in_table"]:
                current["tables"].append(current["_current_table"])
                current["_current_table"] = []
                current["_in_table"] = False

    # Flush last table
    if current and current["_in_table"]:
        current["tables"].append(current["_current_table"])

    return sections


# ── Build DOCX ─────────────────────────────────────────────────────────

def build_docx(sections, out_path):
    doc = Document()
    set_narrow_margins(doc)

    # ── Title ──
    title = doc.add_heading("ReColor — Consolidated Test Cases", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    # Subtitle / instruction
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Mark ✓ under Pass or Fail. Add remarks for failures.")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run.italic = True

    doc.add_paragraph()  # spacer

    for sec in sections:
        if sec["heading"] is None and sec.get("blockquote"):
            continue  # already handled as subtitle

        if sec["heading"]:
            level = min(sec["level"], 4)
            h = doc.add_heading(sec["heading"], level=level)
            # Color the heading
            for run in h.runs:
                if level <= 2:
                    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
                else:
                    run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

        for tbl_data in sec["tables"]:
            if not tbl_data:
                continue

            num_cols = len(tbl_data[0])
            table = doc.add_table(rows=0, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"

            for r_idx, row_cells in enumerate(tbl_data):
                row = table.add_row()
                for c_idx, val in enumerate(row_cells):
                    if c_idx < num_cols:
                        cell = row.cells[c_idx]
                        cell.text = val
                        style_data_cell(cell)

            # Style header
            if len(table.rows) > 0:
                style_header_row(table.rows[0])

            # Zebra stripes
            add_alternating_shading(table)

            # Column widths heuristic based on column count
            if num_cols == 9:
                # TC# | ID | Module | Scenario | Action | Expected | Pass | Fail | Notes
                set_column_widths(table, [1.0, 2.0, 2.8, 3.5, 3.5, 4.0, 1.2, 1.2, 3.5])
            elif num_cols == 3:
                set_column_widths(table, [5.0, 3.0, 2.0])

            doc.add_paragraph()  # spacer after table

    doc.save(out_path)
    print(f"[OK] Saved: {out_path}")


# ── Main ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    sections = parse_md(MD_PATH)
    build_docx(sections, DOCX_PATH)
