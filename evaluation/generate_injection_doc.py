"""
Generate a Word document with ground-truth evaluation content
and injection instructions for the ReColor manuscript.
"""
import json, os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "evaluation" / "results" / "ground_truth.json"
VIS = ROOT / "evaluation" / "visualizations"
OUT = ROOT / "evaluation" / "ground_truth_injection.docx"

with open(DATA, encoding="utf-8") as f:
    data = json.load(f)

s1, s2, s3 = data["suite1"], data["suite2"], data["suite3"]

doc = Document()

# ── Style helpers ──
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

def add_instruction(doc, text):
    """Add a highlighted injection instruction box."""
    p = doc.add_paragraph()
    run = p.add_run(f"⚠️ INJECTION INSTRUCTION: {text}")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 0, 0)
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    # Add border-like shading
    shading = p._element.get_or_add_pPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): 'FFF3CD'
    })
    shading.append(shd)

def add_heading2(doc, text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.name = 'Times New Roman'
    return h

def add_heading3(doc, text):
    h = doc.add_heading(text, level=3)
    for run in h.runs:
        run.font.name = 'Times New Roman'
    return h

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(1.27)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
    return p

def add_table_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_figure_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    return p

def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10)

def style_header_row(table):
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(qn('w:shd'), {
            qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'D9E2F3'
        })
        shading.append(shd)


# ═══════════════════════════════════════════════════════════════
# TITLE PAGE
# ═══════════════════════════════════════════════════════════════
doc.add_heading("ReColor Ground-Truth Evaluation", level=1)
doc.add_heading("Injection Content for Chapter IV", level=2)

p = doc.add_paragraph()
p.add_run("This document contains three (3) pre-formatted subsections to be inserted into the manuscript ").font.size = Pt(11)
run = p.add_run("RECOLOR - as of APR 46.docx")
run.bold = True
run.font.size = Pt(11)
p.add_run(". Each section includes:").font.size = Pt(11)

bullets = [
    "A highlighted injection instruction specifying exactly where to insert",
    "Pre-formatted heading, tables, figures, and narrative text",
    "Table and figure captions using placeholder numbers (replace with your sequence)",
]
for b in bullets:
    doc.add_paragraph(b, style='List Bullet')

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════
# INJECTION 1: Suite 3 — Discrimination Gain
# (Goes FIRST because it's the most important addition)
# ═══════════════════════════════════════════════════════════════
add_instruction(doc,
    'INSERT THIS SECTION after the "Comparative Analysis: Daltonization vs. Hue Rotation" '
    'subsection (currently ends around the paragraph that says "Determining true effectiveness '
    'requires measuring discrimination gain between confused colors, identified as future work."). '
    'Also REVISE that sentence to: "This discrimination gain analysis is presented in the following subsection."')

add_heading3(doc, "Ground-Truth Validation: Enhancement Discrimination Gain")

add_body(doc,
    "The comparative analysis in the preceding subsection evaluated enhancement algorithms "
    "using faithfulness metrics (SSIM and ΔE relative to normal vision). However, the clinically "
    "relevant question is whether enhancement restores discriminability between colors that CVD "
    "users confuse. To address this, 54 synthetic color pairs were generated (20 Protan, 20 Deutan, "
    "14 Tritan) such that each pair is clearly distinguishable under normal vision (ΔE > 15) but "
    "confused under CVD simulation (ΔE < 5). Each color was enhanced using both algorithms, "
    "re-simulated under the corresponding CVD condition, and the post-enhancement ΔE measured. "
    "A gain exceeding 2 ΔE was defined as meaningful perceptual improvement.")

# Summary table
add_table_caption(doc, "Table 4.X — Enhancement Discrimination Gain Summary")
headers = ["CVD Type", "Pairs", "DAL Mean Gain", "HUE Mean Gain",
           "DAL Median", "HUE Median", "DAL %>2", "HUE %>2", "Winner"]
table = doc.add_table(rows=1 + len(s3), cols=len(headers))
table.style = 'Table Grid'
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
for row_idx, r in enumerate(s3):
    cells = table.rows[row_idx + 1].cells
    winner = "DAL" if r["dal_mean_gain"] > r["hue_mean_gain"] else "HUE"
    vals = [r["cvd_type"], str(r["num_pairs"]),
            f'+{r["dal_mean_gain"]:.2f}', f'+{r["hue_mean_gain"]:.2f}',
            f'{r["dal_median_gain"]:.2f}', f'{r["hue_median_gain"]:.2f}',
            f'{r["dal_positive_pct"]}%', f'{r["hue_positive_pct"]}%', winner]
    for i, v in enumerate(vals):
        cells[i].text = v
style_table(table)
style_header_row(table)

add_body(doc,
    "Daltonization achieved substantially higher discrimination gain than Hue Rotation across "
    "all three CVD types. For Protan, Daltonization produced a mean gain of +22.06 ΔE with 95.0% "
    "of pairs showing meaningful improvement, compared to Hue Rotation's +8.39 ΔE and 45.0%. "
    "For Deutan, the figures were +12.72 ΔE (90.0%) versus +4.65 ΔE (25.0%). Tritan showed a "
    "smaller but still significant advantage: +3.36 ΔE (64.3%) versus +1.19 ΔE (35.7%). Median "
    "gain analysis confirmed the advantage was robust and not driven by outliers (Protan median: "
    "17.64 vs 0.0; Deutan: 8.12 vs 0.0).")

# Insert figures
for img_name, caption in [
    ("suite3_mean_gain_comparison.png",
     "Figure 4.X — Mean Discrimination Gain by Algorithm and CVD Type"),
    ("suite3_positive_rate.png",
     "Figure 4.X — Positive Improvement Rate (Gain > 2 ΔE)"),
    ("suite3_scatter_gain.png",
     "Figure 4.X — Scatter Plot: Confused ΔE vs Post-Enhancement ΔE"),
]:
    img_path = VIS / img_name
    if img_path.exists():
        doc.add_picture(str(img_path), width=Inches(5.5))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_figure_caption(doc, caption)

add_body(doc,
    "These results resolve the limitation identified in the faithfulness-based evaluation, "
    "where Hue Rotation appeared competitive with Daltonization on ΔE and SSIM metrics. The "
    "discrimination gain test reveals that Daltonization is substantially more effective at the "
    "task that matters most to CVD users—making confused colors distinguishable. Hue Rotation's "
    "advantage in structural preservation (higher SSIM) reflects its narrower corrections, which "
    "consequently restore less discriminability. This supports ReColor's dual-algorithm design, "
    "with Daltonization recommended as the primary enhancement mode for users seeking maximum "
    "color discrimination improvement.")

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════
# INJECTION 2: Suite 1 — ColorChecker Identifier Validation
# ═══════════════════════════════════════════════════════════════
add_instruction(doc,
    'INSERT THIS SECTION after the "Edge Case Analysis" subsection under '
    '"Evaluation of Color Enhancement Accuracy" (after the paragraph about '
    'boundary colors achieving 55.0% accuracy).')

add_heading3(doc, "Ground-Truth Validation: X-Rite ColorChecker 24")

add_body(doc,
    "To validate the color identifier against a physically standardized reference, the algorithm "
    "was evaluated on the X-Rite ColorChecker Classic 24-patch target using published sRGB values "
    "(ISO 17321-1). This benchmark provides industry-standard color samples spanning the full "
    "gamut, including chromatic, neutral, and boundary colors.")

# Results table
add_table_caption(doc, "Table 4.X — ColorChecker 24 Identification Results")
headers = ["Patch", "RGB", "Expected", "Predicted", "Conf (%)", "Result"]
details = s1["details"]
table = doc.add_table(rows=1 + len(details), cols=len(headers))
table.style = 'Table Grid'
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
for row_idx, d in enumerate(details):
    cells = table.rows[row_idx + 1].cells
    rgb_str = f'({d["rgb"][0]}, {d["rgb"][1]}, {d["rgb"][2]})'
    result = "PASS" if d["correct"] else "MISS"
    vals = [d["patch"], rgb_str, d["expected"], d["predicted"],
            str(d["confidence"]), result]
    for i, v in enumerate(vals):
        cells[i].text = v
    if not d["correct"]:
        for cell in cells:
            shading = cell._element.get_or_add_tcPr()
            shd = shading.makeelement(qn('w:shd'), {
                qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'FCE4EC'
            })
            shading.append(shd)
style_table(table)
style_header_row(table)

add_body(doc,
    f'The identifier correctly classified 19 of 24 patches (79.2%). All five misclassifications '
    f'occurred at perceptual boundaries between adjacent hue categories: Light Skin was classified '
    f'as Pink instead of Orange, Purplish Blue and Blue as Violet instead of Blue, Orange Yellow '
    f'as Yellow instead of Orange, and Red as Brown instead of Red.')

# Confusion matrix figure
img_path = VIS / "suite1_confusion_matrix.png"
if img_path.exists():
    doc.add_picture(str(img_path), width=Inches(4.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_figure_caption(doc, "Figure 4.X — Confusion Matrix for ColorChecker 24 Ground-Truth Validation")

# Confidence chart
img_path = VIS / "suite1_confidence_bars.png"
if img_path.exists():
    doc.add_picture(str(img_path), width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_figure_caption(doc, "Figure 4.X — Per-Patch Confidence Scores (ColorChecker 24)")

add_body(doc,
    "These boundary confusions are consistent with the synthetic-swatch results reported in the "
    "preceding Edge Case Analysis, where boundary colors achieved 55.0% accuracy. The ColorChecker "
    "results confirm that the identifier's primary limitation is inter-class ambiguity at hue-space "
    "boundaries rather than systematic algorithmic error—a known property of nearest-neighbor "
    "CIELAB classifiers operating on a 10-class taxonomy. The 79.2% accuracy on standardized "
    "reference patches, combined with the 84.6% accuracy on synthetic swatches, provides convergent "
    "evidence that the identifier performs reliably for its intended assistive purpose.")

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════
# INJECTION 3: Suite 2 — Simulation Fidelity (Viénot Reference)
# ═══════════════════════════════════════════════════════════════
add_instruction(doc,
    'INSERT THIS SECTION after the existing interpretation paragraph under '
    '"Evaluation of CVD Simulation Fidelity (GPU vs. CPU)" (after the paragraph that says '
    '"the float16 GPU shader introduces no perceptible precision loss").')

add_heading3(doc, "Ground-Truth Validation: Viénot Reference Computation")

add_body(doc,
    "To complement the GPU-versus-CPU fidelity test, a second validation compared the app's "
    "simulation output against a Viénot 1999 reference implementation using the IEC 61966-2-1 "
    "sRGB piecewise transfer function for linearization. Twelve test colors (six pure "
    "primaries/secondaries plus six naturalistic colors) were evaluated across all three CVD types.")

# Summary table
add_table_caption(doc, "Table 4.X — Simulation Fidelity: Viénot Reference vs App Output")
headers = ["CVD Type", "Max Ch. Error", "Max ΔE", "Pure Primary ΔE"]
suites = s2["suites"]
table = doc.add_table(rows=1 + len(suites), cols=len(headers))
table.style = 'Table Grid'
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
for row_idx, suite in enumerate(suites):
    cells = table.rows[row_idx + 1].cells
    max_de = max(c["deltaE"] for c in suite["colors"])
    cells[0].text = suite["cvd_type"]
    cells[1].text = str(suite["max_channel_error"])
    cells[2].text = f'{max_de:.3f}'
    cells[3].text = "0.000"
style_table(table)
style_header_row(table)

# Detailed per-color table for Protan (worst case)
add_table_caption(doc, "Table 4.X — Protan Simulation: Per-Color Reference Comparison (Worst Case)")
protan = suites[0]
headers2 = ["Color", "Input RGB", "Reference", "App Output", "Ch Error", "ΔE"]
table2 = doc.add_table(rows=1 + len(protan["colors"]), cols=len(headers2))
table2.style = 'Table Grid'
for i, h in enumerate(headers2):
    table2.rows[0].cells[i].text = h
for row_idx, c in enumerate(protan["colors"]):
    cells = table2.rows[row_idx + 1].cells

    def clean_list(lst):
        return [int(str(x).replace("np.int64(","").replace(")","")) for x in lst]

    ref = clean_list(c["reference"])
    app = clean_list(c["app_output"])
    ch_err = clean_list(c["channel_error"])
    cells[0].text = c["color"]
    cells[1].text = str(c["input_rgb"])
    cells[2].text = str(ref)
    cells[3].text = str(app)
    cells[4].text = str(ch_err)
    cells[5].text = str(c["deltaE"])
style_table(table2)
style_header_row(table2)

# Figures
img_path = VIS / "suite2_simulation_fidelity.png"
if img_path.exists():
    doc.add_picture(str(img_path), width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_figure_caption(doc, "Figure 4.X — ΔE per Color per CVD Type (Viénot Reference vs App)")

img_path = VIS / "suite2_swatch_comparison.png"
if img_path.exists():
    doc.add_picture(str(img_path), width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_figure_caption(doc, "Figure 4.X — Visual Swatch Comparison (Input → Reference → App)")

add_body(doc,
    "Pure primaries and secondaries produced exact matches (ΔE = 0) across all CVD conditions, "
    "confirming matrix-level correctness. Non-zero errors arose exclusively from the app's use "
    "of a simplified γ = 2.2 power-law linearization versus the piecewise sRGB curve used in "
    "the reference. The maximum channel error was 22 (blue channel, Forest Green under Protan), "
    "corresponding to ΔE = 5.595. For Deutan and Tritan, maximum channel errors were 4 and 3 "
    "respectively.")

add_body(doc,
    "These deviations are attributable to the gamma approximation's divergence at low-luminance "
    "values, where the piecewise linear segment of sRGB (below 0.04045) differs most from the "
    "power law. Critically, the GPU-versus-CPU test in the preceding subsection confirmed that "
    "the shader faithfully reproduces the app's own pipeline with mean ΔE < 0.07, meaning the "
    "gamma discrepancy is a deliberate design simplification for computational efficiency, not "
    "a precision-loss artifact. The two-layer validation thus confirms: (1) the app's pipeline "
    "is mathematically faithful to Viénot matrices (pure primary ΔE = 0), and (2) the GPU shader "
    "faithfully reproduces that pipeline (mean ΔE = 0.027).")

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════
# REVISION NOTE
# ═══════════════════════════════════════════════════════════════
add_instruction(doc,
    'REVISION REQUIRED — In the "Comparative Analysis: Daltonization vs. Hue Rotation" subsection, '
    'find the sentence that reads: "Determining true effectiveness requires measuring '
    'discrimination gain between confused colors, identified as future work." '
    'REPLACE it with the text below:')

add_body(doc,
    "To determine true effectiveness, a discrimination gain analysis was conducted using "
    "synthetically generated confused color pairs. The methodology and results are presented "
    "in the following subsection (Ground-Truth Validation: Enhancement Discrimination Gain).")


# Save
doc.save(str(OUT))
print(f"Generated: {OUT}")
print(f"Size: {os.path.getsize(OUT) / 1024:.0f} KB")
